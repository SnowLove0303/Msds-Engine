"""MSDS 模板覆写工作台（识别模块同款 UI）.

Chrome 与交互照搬识别模块 MsdsViewer：顶部模板切换 + 文件名 + 检索 +
状态；左侧章节导航；右侧三页签（覆写编辑 / 原版对照 / 格式与审计）；
结构化网格按源列宽展示、首行高亮、单元格只读标签 + 可编辑值。

与识别模块的区别（覆写专属）：
- 固定导入内嵌 CN/EN 冠志模板（工作副本编辑，源模板只读 + 哈希保护）；
- 值格编辑、行新增/删除、自动连续重排、导出前审计拦截；
- 受控标签编辑：仅 Section 9，经确认开关后可用，其余标签永久锁定。

所有写入经 msds_template_editor 核心 API，本文件只负责界面，不碰格式基因。
"""

from __future__ import annotations

import json
import shutil
import tempfile
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from pathlib import Path

from docx import Document

from msds_template_editor import (
    EMBEDDED_TEMPLATE_CN,
    EMBEDDED_TEMPLATE_EN,
    MutationViolation,
    audit_document,
    cell_views,
    create_work_copy,
    file_sha256,
    is_row_editable,
    renumber_document,
    row_cells_with_spans,
    row_prefix,
    safe_add_row_after,
    safe_delete_row,
    set_field_value,
    unique_cells,
    visible_text,
)

FIXED_TEMPLATES = {
    "CN 冠志": EMBEDDED_TEMPLATE_CN,
    "EN 冠志": EMBEDDED_TEMPLATE_EN,
}

SECTION_TITLES = [
    "物料及供应商标识", "危险性概述", "成分/组成资料", "急救措施",
    "消防措施", "意外泄漏措施", "操作和储存", "接触控制/个人防护",
    "物理和化学特性", "稳定性和反应性", "毒性资料", "生态信息",
    "处理注意事项", "运输信息", "法规信息", "其他信息",
]


def table_title(table, fallback: str) -> str:
    try:
        text = visible_text(table.rows[0].cells[0]).strip()
    except IndexError:
        text = ""
    return text or fallback


class OverwriteStudio:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("MSDS 模板覆写")
        self.root.geometry("1280x820")
        self.root.minsize(900, 620)
        self.document = None
        self.source_path: Path | None = None
        self.work_path: Path | None = None
        self.work_dir = Path(tempfile.mkdtemp(prefix="msds-overwrite-"))
        self.source_hash = ""
        self.dirty = False
        self.template_name = "CN 冠志"
        self.current_table = 0
        self.query = tk.StringVar()
        self.allow_label_edit = tk.BooleanVar(value=False)
        self.visible_tables: list[int] = []
        self._text_widgets: list[tuple] = []
        self._baseline_snapshot: list[dict] = []
        self._build_ui()
        self.load_template("CN 冠志")

    # ---------- 界面骨架（同识别模块） ----------
    def _build_ui(self):
        self.root.protocol("WM_DELETE_WINDOW", self.close)
        top = ttk.Frame(self.root, padding=(10, 8))
        top.pack(fill="x")
        self.cn_button = ttk.Button(top, text="CN 冠志", command=lambda: self.switch_template("CN 冠志"))
        self.cn_button.pack(side="left")
        self.en_button = ttk.Button(top, text="EN 冠志", command=lambda: self.switch_template("EN 冠志"))
        self.en_button.pack(side="left", padx=(6, 0))
        self.export_button = ttk.Button(top, text="导出编辑后 DOCX", command=self.export)
        self.export_button.pack(side="left", padx=(6, 0))
        self.audit_button = ttk.Button(top, text="审计", command=self.run_audit)
        self.audit_button.pack(side="left", padx=(6, 0))
        self.filename = ttk.Label(top, text="尚未载入模板")
        self.filename.pack(side="left", padx=12)
        ttk.Label(top, text="检索：").pack(side="left", padx=(18, 4))
        entry = ttk.Entry(top, textvariable=self.query, width=34)
        entry.pack(side="left")
        self.query.trace_add("write", lambda *_: self.refresh_list())
        self.status = ttk.Label(top, text="只写值格，不改模板")
        self.status.pack(side="right")

        body = ttk.PanedWindow(self.root, orient="horizontal")
        body.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        nav = ttk.Frame(body, width=330)
        body.add(nav, weight=1)
        ttk.Label(nav, text="章节 / 表格 / 搜索结果").pack(anchor="w", pady=(0, 5))
        self.listbox = tk.Listbox(nav, exportselection=False, activestyle="dotbox")
        list_scroll = ttk.Scrollbar(nav, orient="vertical", command=self.listbox.yview)
        self.listbox.configure(yscrollcommand=list_scroll.set)
        self.listbox.pack(side="left", fill="both", expand=True)
        list_scroll.pack(side="right", fill="y")
        self.listbox.bind("<<ListboxSelect>>", self._select_section)

        viewer = ttk.Frame(body)
        body.add(viewer, weight=4)
        self.tabs = ttk.Notebook(viewer)
        self.tabs.pack(fill="both", expand=True)
        self.edit_tab = ttk.Frame(self.tabs)
        self.base_tab = ttk.Frame(self.tabs)
        self.format_tab = ttk.Frame(self.tabs)
        self.tabs.add(self.edit_tab, text="覆写编辑（默认）")
        self.tabs.add(self.base_tab, text="原版对照")
        self.tabs.add(self.format_tab, text="格式与审计")

        self.canvas = tk.Canvas(self.edit_tab, background="white", highlightthickness=0)
        vertical = ttk.Scrollbar(self.edit_tab, orient="vertical", command=self.canvas.yview)
        horizontal = ttk.Scrollbar(self.edit_tab, orient="horizontal", command=self.canvas.xview)
        self.canvas.configure(yscrollcommand=vertical.set, xscrollcommand=horizontal.set)
        self.canvas.grid(row=0, column=0, sticky="nsew")
        vertical.grid(row=0, column=1, sticky="ns")
        horizontal.grid(row=1, column=0, sticky="ew")
        self.edit_tab.rowconfigure(0, weight=1)
        self.edit_tab.columnconfigure(0, weight=1)
        self.form = ttk.Frame(self.canvas, padding=14)
        self.canvas_window = self.canvas.create_window((0, 0), window=self.form, anchor="nw")
        self.form.bind("<Configure>", lambda _: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.bind("<Configure>", lambda event: self.canvas.itemconfigure(self.canvas_window, width=event.width))
        self.canvas.bind_all("<MouseWheel>", self._mousewheel)

        self.base_canvas = tk.Canvas(self.base_tab, background="#d9dde3", highlightthickness=0)
        base_scroll = ttk.Scrollbar(self.base_tab, orient="vertical", command=self.base_canvas.yview)
        self.base_canvas.configure(yscrollcommand=base_scroll.set)
        self.base_canvas.pack(side="left", fill="both", expand=True)
        base_scroll.pack(side="right", fill="y")
        self.base_content = ttk.Frame(self.base_canvas)
        self.base_window = self.base_canvas.create_window((0, 0), window=self.base_content, anchor="n")
        self.base_content.bind("<Configure>", lambda _: self.base_canvas.configure(scrollregion=self.base_canvas.bbox("all")))
        self.base_canvas.bind("<Configure>", lambda event: self.base_canvas.itemconfigure(self.base_window, width=event.width))

        self.format_view = tk.Text(self.format_tab, wrap="none", font=("Consolas", 9))
        format_y = ttk.Scrollbar(self.format_tab, orient="vertical", command=self.format_view.yview)
        format_x = ttk.Scrollbar(self.format_tab, orient="horizontal", command=self.format_view.xview)
        self.format_view.configure(yscrollcommand=format_y.set, xscrollcommand=format_x.set, state="disabled")
        self.format_view.grid(row=0, column=0, sticky="nsew")
        format_y.grid(row=0, column=1, sticky="ns")
        format_x.grid(row=1, column=0, sticky="ew")
        self.format_tab.rowconfigure(0, weight=1)
        self.format_tab.columnconfigure(0, weight=1)

    def _mousewheel(self, event):
        if self.document is not None:
            self.canvas.yview_scroll(int(-event.delta / 120), "units")

    # ---------- 模板装载（固定 CN/EN） ----------
    def switch_template(self, name: str):
        if name == self.template_name and self.document is not None:
            return
        if self.dirty and not messagebox.askyesno("未导出修改", "当前工作副本有未导出修改，切换模板会丢弃这些修改。继续吗？"):
            return
        self.load_template(name)

    def load_template(self, name: str):
        path = FIXED_TEMPLATES[name].resolve()
        if not path.is_file():
            messagebox.showerror("模板不存在", f"找不到内嵌模板：\n{path}")
            return
        try:
            work_path = create_work_copy(path, self.work_dir)
            document = Document(str(work_path))
            baseline = Document(str(path))
        except Exception as exc:
            messagebox.showerror("模板无法读取", f"无法读取 DOCX：\n{exc}")
            return
        self.document = document
        self.source_path = path
        self.work_path = work_path
        self.source_hash = file_sha256(path)
        self.template_name = name
        self.current_table = 0
        self.allow_label_edit.set(False)
        self._baseline_snapshot = [
            {"title": table_title(table, f"表格 {index + 1}"), "rows": len(table.rows)}
            for index, table in enumerate(baseline.tables)
        ]
        self._set_dirty(False)
        self.filename.configure(text=f"{path.name}（工作副本）")
        self.status.configure(text=f"已载入{name}，源模板只读")
        self.refresh_list()
        self._render_baseline()
        self._render_format()

    # ---------- 左侧导航（同识别模块检索） ----------
    def refresh_list(self):
        self.listbox.delete(0, "end")
        self.visible_tables = []
        if self.document is None:
            return
        query = self.query.get().strip().casefold()
        for index, table in enumerate(self.document.tables):
            title = table_title(table, f"表格 {index + 1}")
            haystack = f"{index + 1:02d} {title}".casefold()
            if query and query not in haystack:
                continue
            self.visible_tables.append(index)
            self.listbox.insert("end", f"{index + 1:02d} {title} · {len(table.rows)} 行")
        self.status.configure(
            text=f"{len(self.visible_tables)} 节" + (f"匹配“{self.query.get().strip()}”" if query else "")
        )
        if self.visible_tables:
            self.listbox.selection_set(0)
            self._select_section()

    def _select_section(self, _event=None):
        selection = self.listbox.curselection()
        if not selection or not self.visible_tables:
            return
        self.current_table = self.visible_tables[selection[0]]
        self._render_current_section()

    # ---------- 中间编辑区 ----------
    def _render_current_section(self):
        for widget in self.form.winfo_children():
            widget.destroy()
        self._text_widgets.clear()
        if self.document is None or self.current_table >= len(self.document.tables):
            return
        table = self.document.tables[self.current_table]
        title = table_title(table, f"表格 {self.current_table + 1}")
        ttk.Label(self.form, text=f"{title}   ·   {len(table.rows)} 行", font=("Microsoft YaHei UI", 13, "bold")).grid(
            row=0, column=0, columnspan=len(table.columns) + 1, sticky="w", pady=(0, 10)
        )
        column_count = max(1, len(table.columns))
        for column in range(column_count):
            try:
                width = table.columns[column].width
                minsize = max(130, min(420, int(width.inches * 96))) if width else 220
            except (AttributeError, TypeError, ValueError):
                minsize = 220
            self.form.columnconfigure(column + 1, weight=0, minsize=minsize)
        self.form.columnconfigure(0, weight=0, minsize=88)
        for row_index, row in enumerate(table.rows):
            self._render_row(table, row_index, row, column_count)

    def _render_row(self, table, row_index: int, row, column_count: int):
        header = row_index == 0
        background = "#f5f8fc" if header else "white"
        action = tk.Frame(self.form, background="#f8fafc", width=88)
        action.grid(row=row_index + 1, column=0, sticky="nsew", padx=(0, 1), pady=(0, 1))
        ttk.Label(action, text=f"行 {row_index + 1}", foreground="#667085").pack(pady=(4, 0))
        sequence = row_prefix(row)
        if sequence:
            ttk.Label(action, text=f"序 {sequence}", foreground="#98a2b3", font=("Microsoft YaHei UI", 8)).pack()
        add_btn = ttk.Button(action, text="＋", width=6, command=lambda i=row_index: self.add_row(i))
        add_btn.pack(pady=1)
        if row_index == 0:
            add_btn.configure(state="disabled")
        if row_index > 0 and is_row_editable(row) and len(table.rows) > 2:
            ttk.Button(action, text="×", width=6, command=lambda i=row_index: self.remove_row(i)).pack(pady=1)
        views = cell_views(self.current_table, row_index, row)
        covered = set()
        for view in views:
            for column in range(view.start_column, view.start_column + view.span):
                covered.add(column)
            cell_frame = tk.Frame(self.form, background=background, relief="solid", borderwidth=1)
            cell_frame.grid(row=row_index + 1, column=view.start_column + 1, columnspan=view.span,
                            sticky="nsew", padx=(0, 1), pady=(0, 1))
            self._render_cell(cell_frame, view, background)
        for column in range(column_count):
            if column not in covered:
                tk.Frame(self.form, background=background, relief="solid", borderwidth=1).grid(
                    row=row_index + 1, column=column + 1, sticky="nsew", padx=(0, 1), pady=(0, 1)
                )

    def _render_cell(self, parent, view, background: str):
        if view.kind == "note":
            ttk.Label(parent, text="说明行 · 单列值", foreground="#b54708", font=("Microsoft YaHei UI", 8)).pack(
                anchor="e", padx=5, pady=(3, 0))
        if view.label_field:
            if self._label_editable_here():
                self._render_edit(parent, view.label_field, bold=True, background=background)
            else:
                font = ("Microsoft YaHei UI", 10, "bold")
                tk.Label(parent, text=view.label_field.value, font=font, background=background,
                         foreground="#182230", anchor="w", justify="left", wraplength=700).pack(
                    fill="x", padx=5, pady=2)
        if view.value_field:
            self._render_edit(parent, view.value_field, bold=False, background=background)
        if not view.label_field and not view.value_field:
            ttk.Label(parent, text="结构单元格", foreground="#98a2b3").pack(padx=5, pady=2)

    def _label_editable_here(self) -> bool:
        # 受控标签编辑：仅 Section 9（第 9 张表），且开关已确认开启。
        return bool(self.allow_label_edit.get()) and self.current_table == 8

    def _render_edit(self, parent, field, bold: bool, background: str):
        line = tk.Frame(parent, background=background)
        line.pack(fill="x", padx=5, pady=2)
        font = ("Microsoft YaHei UI", 10, "bold") if bold else ("Microsoft YaHei UI", 10)
        if "\n" in field.value or len(field.value) > 75:
            text = tk.Text(line, height=min(5, max(2, field.value.count("\n") + 1)), wrap="word", undo=True,
                           font=font, relief="flat", borderwidth=0, highlightthickness=0, background="white")
            text.insert("1.0", field.value)
            text.pack(side="left", fill="x", expand=True)
            text.bind("<KeyRelease>", lambda _e, w=text, f=field: self._text_changed(w, f))
            text.bind("<FocusOut>", lambda _e, w=text, f=field: self._text_changed(w, f))
            self._text_widgets.append((text, field))
        else:
            value = tk.StringVar(value=field.value)
            entry = tk.Entry(line, textvariable=value, font=font, relief="flat", borderwidth=0,
                             highlightthickness=0, background="white")
            entry.pack(side="left", fill="x", expand=True)
            value.trace_add("write", lambda *_args, v=value, f=field: self._value_changed(v.get(), f))

    # ---------- 写入（经核心 API） ----------
    def _value_changed(self, value, field):
        if self.document is not None and value != field.value:
            try:
                set_field_value(field, value, allow_label_edit=False)
                self._set_dirty()
            except MutationViolation as exc:
                messagebox.showerror("写入约束拦截", str(exc))

    def _text_changed(self, widget, field):
        value = widget.get("1.0", "end-1c")
        if self.document is not None and value != field.value:
            try:
                set_field_value(field, value, allow_label_edit=self._label_editable_here())
                self._set_dirty()
            except MutationViolation as exc:
                messagebox.showerror("写入约束拦截", str(exc))

    def _sync_texts(self):
        for widget, field in self._text_widgets:
            self._text_changed(widget, field)

    def toggle_label_edit(self):
        if self.allow_label_edit.get():
            confirmed = messagebox.askyesno(
                "受控标签编辑确认",
                "仅 Section 9（理化特性）允许修改标签，用于补充测定条件或限定词，\n"
                "序号前缀与全部格式保持锁定。确认开启？",
            )
            if not confirmed:
                self.allow_label_edit.set(False)
                return
        self._render_current_section()

    def add_row(self, row_index: int):
        if self.document is None:
            return
        self._sync_texts()
        try:
            safe_add_row_after(self.document.tables[self.current_table], row_index, auto_renumber=True)
        except MutationViolation as exc:
            messagebox.showerror("新增被拦截", str(exc))
            return
        self._set_dirty()
        self.refresh_list()
        self._render_format()

    def remove_row(self, row_index: int):
        if self.document is None:
            return
        table = self.document.tables[self.current_table]
        if not is_row_editable(table.rows[row_index]):
            messagebox.showwarning("无法删除", "该行是受保护的表头或结构行，无法删除。")
            return
        if not messagebox.askyesno("确认删除", f"确认删除第 {row_index + 1} 行？后续序号自动连续。"):
            return
        self._sync_texts()
        try:
            safe_delete_row(table, row_index, auto_renumber=True)
        except MutationViolation as exc:
            messagebox.showerror("删除被拦截", str(exc))
            return
        self._set_dirty()
        self.refresh_list()
        self._render_format()

    # ---------- 原版对照（只读） ----------
    def _render_baseline(self):
        for widget in self.base_content.winfo_children():
            widget.destroy()
        if not self._baseline_snapshot:
            return
        ttk.Label(self.base_content, text=f"{self.template_name} 内嵌模板只读对照（{self.source_hash[:12]}…）",
                  font=("Microsoft YaHei UI", 13, "bold")).pack(pady=10)
        ttk.Label(self.base_content, text="对照基线只读，编辑区改动不影响本页；源模板哈希变更即停工。",
                  foreground="#667085").pack(pady=(0, 10))
        for index, item in enumerate(self._baseline_snapshot):
            ttk.Label(self.base_content, text=f"{index + 1:02d} {item['title']} · {item['rows']} 行").pack()

    # ---------- 格式与审计 ----------
    def _render_format(self):
        if self.document is None:
            return
        errors = audit_document(self.document)
        payload = {
            "template": self.template_name,
            "source_hash": self.source_hash,
            "tables": [
                {"section": index + 1, "rows": len(table.rows), "columns": len(table.columns)}
                for index, table in enumerate(self.document.tables)
            ],
            "audit_errors": errors,
            "dirty": self.dirty,
        }
        self.format_view.configure(state="normal")
        self.format_view.delete("1.0", "end")
        self.format_view.insert("1.0", json.dumps(payload, ensure_ascii=False, indent=2))
        self.format_view.configure(state="disabled")

    def run_audit(self):
        if self.document is None:
            return
        self._sync_texts()
        errors = audit_document(self.document)
        self._render_format()
        if errors:
            messagebox.showerror("审计未通过", "\n".join(errors[:10]))
            self.status.configure(text=f"审计发现 {len(errors)} 个问题")
        else:
            messagebox.showinfo("审计通过", "16 节序号连续，无违规加粗。")
            self.status.configure(text="审计通过")

    def export(self):
        if self.document is None or self.source_path is None:
            return
        self._sync_texts()
        renumber_errors = renumber_document(self.document)
        _ = renumber_errors
        errors = audit_document(self.document)
        self._render_format()
        if errors:
            messagebox.showerror("导出已阻止", f"请先处理 {len(errors)} 个审计问题。")
            return
        initial = f"{self.source_path.stem}_编辑后.docx"
        path = filedialog.asksaveasfilename(title="导出编辑后 DOCX", initialdir=str(self.source_path.parent),
                                            initialfile=initial, defaultextension=".docx",
                                            filetypes=[("Word 文档", "*.docx")])
        if not path:
            return
        output = Path(path).resolve()
        if output == self.source_path:
            messagebox.showerror("禁止覆盖源模板", "请另存为新文件，源模板不会被覆盖。")
            return
        try:
            if self.work_path is None:
                raise RuntimeError("未建立模板工作副本")
            if file_sha256(self.source_path) != self.source_hash:
                raise RuntimeError("源模板哈希已变化，停止导出以保护基线。")
            self.document.save(str(self.work_path))
            shutil.copy2(self.work_path, output)
            self._set_dirty(False)
            messagebox.showinfo("导出成功", f"文件已导出至：\n{output}")
        except Exception as exc:
            messagebox.showerror("导出失败", f"源模板未被更改。\n\n{exc}")

    def _set_dirty(self, dirty: bool = True):
        self.dirty = dirty
        self.status.configure(text="● 有未导出修改" if dirty else "工作副本干净")

    def confirm_discard(self) -> bool:
        if not self.dirty:
            return True
        return messagebox.askyesno("未导出修改", "当前有未导出修改，确认放弃并继续？")

    def close(self):
        if self.confirm_discard():
            shutil.rmtree(self.work_dir, ignore_errors=True)
            self.root.destroy()


def build_label_switch(parent, studio: OverwriteStudio):
    """受控标签编辑开关（仅 Section 9）：复选框 + 确认，与编辑区联动。"""
    switch = ttk.Checkbutton(parent, text="受控：允许修改标签（仅 Section 9）",
                             variable=studio.allow_label_edit, command=studio.toggle_label_edit)
    switch.pack(side="left", padx=(12, 0))
    return switch


def main(argv: list[str] | None = None) -> int:
    root = tk.Tk()
    app = OverwriteStudio(root)
    build_label_switch(root.winfo_children()[0], app)
    root.mainloop()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
