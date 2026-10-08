"""Regression check suite for the MSDS template editor core and write-time constraints."""

from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from msds_template_editor import (
    DEFAULT_TEMPLATE,
    EMBEDDED_TEMPLATE_CN,
    EMBEDDED_TEMPLATE_EN,
    TEMPLATE_LIBRARY,
    MutationViolation,
    add_note_row_after,
    build_export_docx_name,
    add_row_after,
    move_row,
    audit_document,
    cell_views,
    create_work_copy,
    delete_row,
    file_sha256,
    normalize_value_text,
    renumber_document,
    renumber_table,
    row_fields,
    row_prefix,
    safe_delete_row,
    set_field_value,
    set_label_value,
    write_cell_value,
    write_section9_property,
)


FORMAT_TAGS = {qn(name) for name in ("w:pPr", "w:tblPr", "w:trPr", "w:tcPr", "w:rPr")}


def format_signature(document):
    return tuple(element.xml for element in document.element.body.iter()
                 if element.tag in FORMAT_TAGS)


def main() -> None:
    assert DEFAULT_TEMPLATE.exists(), DEFAULT_TEMPLATE
    assert TEMPLATE_LIBRARY["EN 冠志"].exists(), TEMPLATE_LIBRARY["EN 冠志"]
    # 基线即内嵌模板：不再依赖外部 F:\App Location 源模板目录。
    assert EMBEDDED_TEMPLATE_CN.exists(), EMBEDDED_TEMPLATE_CN
    assert EMBEDDED_TEMPLATE_EN.exists(), EMBEDDED_TEMPLATE_EN
    original_hash = file_sha256(DEFAULT_TEMPLATE)
    document = Document(str(DEFAULT_TEMPLATE))
    assert len(document.tables) == 16
    baseline_shape = [(len(table.rows), len(table.columns)) for table in document.tables]
    baseline_format = format_signature(document)
    baseline_bold_runs = sum(
        1 for table in document.tables for row in table.rows for cell in row.cells
        for paragraph in cell.paragraphs for run in paragraph.runs if run.bold
    )

    # 1. Non-bold value slot editing test
    editable = None
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            fields = row_fields(table_index, row_index, row)
            if fields:
                editable = (table_index, row_index, fields[0])
                break
        if editable:
            break
    assert editable is not None, "template has no non-bold value slot"
    table_index, row_index, field = editable
    assert field.value, "source content must be visible in the editable model"
    set_field_value(field, "测试值")

    # 2. Three-cell row views test (label, sublabel, value)
    three_cell_views = cell_views(10, 3, document.tables[10].rows[3])
    assert len(three_cell_views) == 3
    assert three_cell_views[0].label_field and three_cell_views[1].label_field
    assert three_cell_views[2].value_field
    
    # 3. Label editing toggle test (blocked by default, permitted when enabled)
    try:
        set_field_value(three_cell_views[0].label_field, "自定义标签", allow_label_edit=False)
        assert False, "Should have raised MutationViolation when allow_label_edit=False"
    except MutationViolation:
        pass  # Expected protection
    set_field_value(three_cell_views[0].label_field, "自定义标签", allow_label_edit=True)
    set_field_value(three_cell_views[2].value_field, "自定义值")
    assert cell_views(10, 1, document.tables[10].rows[1])[0].kind == "note"

    # 4. Section 9 label qualifier update test (preserves 9.n prefix and bold styling)
    sec9_cell = document.tables[8].rows[3].cells[0]
    sec9_view = cell_views(8, 3, document.tables[8].rows[3])[0]
    set_label_value(sec9_view, "pH值（原液）：", allow_label_edit=True)
    assert "9.3  pH值（原液）：" in sec9_cell.text, f"Unexpected text: {sec9_cell.text}"

    # 5. Row deletion protection test (Row 0 is protected)
    try:
        safe_delete_row(document.tables[8], 0)
        assert False, "Should have blocked Row 0 deletion"
    except MutationViolation:
        pass

    # 6. Row addition and deletion test
    table = document.tables[table_index]
    before_rows = len(table.rows)
    add_row_after(table, row_index)
    assert len(table.rows) == before_rows + 1
    renumber_document(document)
    delete_row(table, row_index + 1)
    assert len(table.rows) == before_rows

    # 7. Section 9 row deletion and continuous renumbering test
    sec9_doc = Document(str(DEFAULT_TEMPLATE))
    sec9_tbl = sec9_doc.tables[8]
    assert "9.5" in sec9_tbl.rows[5].cells[0].text
    delete_row(sec9_tbl, 5, auto_renumber=True)  # Delete 9.5 初沸点
    assert "9.5  闪点：" in sec9_tbl.rows[5].cells[0].text, f"Expected 9.5, got: {sec9_tbl.rows[5].cells[0].text}"
    assert "9.22 其他信息：" in sec9_tbl.rows[22].cells[0].text, f"Expected 9.22, got: {sec9_tbl.rows[22].cells[0].text}"

    # 8. Section 1 (Table 0) renumbering integrity test
    t0 = sec9_doc.tables[0]
    renumber_table(t0, 1)
    assert t0.rows[0].cells[0].text.strip() == "1.物料及供应商标识", "Row 0 title must not be corrupted"
    assert "1.1  产品名称：" in t0.rows[1].cells[0].text
    assert "1.2  产品使用建议和使用限制：" in t0.rows[4].cells[0].text
    assert "1.3  供应商信息：" in t0.rows[5].cells[0].text

    # 9. Continuous renumbering test after row addition
    numbering_doc = Document(str(DEFAULT_TEMPLATE))
    numbering_table = numbering_doc.tables[8]
    before_prefixes = [row_prefix(row) for row in numbering_table.rows[1:4]]
    add_row_after(numbering_table, 1)
    renumber_document(numbering_doc)
    after_prefixes = [row_prefix(row) for row in numbering_table.rows[1:5]]
    assert before_prefixes == ["9.1", "9.2", "9.3"]
    assert after_prefixes == ["9.1", "9.2", "9.3", "9.4"]

    # 10. Write cell value & whitespace normalization test
    test_cell = numbering_doc.tables[0].rows[1].cells[1]
    write_cell_value(test_cell, "  标准化值写入\r\n\n  ")
    assert test_cell.text == "标准化值写入"

    # 11. Section 9 high-level property writer test (write_section9_property)
    sec9_prop_tbl = numbering_doc.tables[8]
    write_section9_property(sec9_prop_tbl, "闪点", "> 100 ℃", new_label="闪点（闭杯）：", allow_label_edit=True)
    matched_row = [r for r in sec9_prop_tbl.rows if "闪点" in r.cells[0].text][0]
    assert "闪点（闭杯）：" in matched_row.cells[0].text
    assert matched_row.cells[-1].text == "> 100 ℃"

    # 12. Document audit test (audit_document)
    renumber_document(numbering_doc)
    audit_errs = audit_document(numbering_doc)
    assert not audit_errs, f"Audit errors found: {audit_errs}"

    # 13. Export, reopened signature equality, and audit verification
    with tempfile.TemporaryDirectory(prefix="msds-editor-") as temp_dir:
        work_dir = Path(temp_dir) / "work"
        en_copy = create_work_copy(TEMPLATE_LIBRARY["EN 冠志"], work_dir)
        assert en_copy.exists()
        assert len(Document(str(en_copy)).tables) > 0
        assert file_sha256(TEMPLATE_LIBRARY["EN 冠志"]) == file_sha256(en_copy)
        output = Path(temp_dir) / "export.docx"
        document.save(str(output))
        reopened = Document(str(output))
        assert [(len(table.rows), len(table.columns)) for table in reopened.tables] == baseline_shape
        assert format_signature(reopened) == baseline_format, "Format tags XML signature must match baseline"
        assert sum(
            1 for table in reopened.tables for row in table.rows for cell in row.cells
            for paragraph in cell.paragraphs for run in paragraph.runs if run.bold
        ) == baseline_bold_runs
        all_text = "\n".join(
            cell.text for table in reopened.tables for row in table.rows for cell in row.cells
        )
        assert "自定义标签" in all_text
        assert "自定义值" in all_text
        assert "测试值" in all_text
    assert file_sha256(DEFAULT_TEMPLATE) == original_hash

    # 14. Note row creation and row movement test
    sec11_doc = Document(str(DEFAULT_TEMPLATE))
    sec11_tbl = sec11_doc.tables[10]
    orig_s11_rows = len(sec11_tbl.rows)
    add_note_row_after(sec11_tbl, 0, "【测试节顶部说明行】")
    assert len(sec11_tbl.rows) == orig_s11_rows + 1
    assert "【测试节顶部说明行】" in sec11_tbl.rows[1].cells[0].text

    # Move note row down
    move_row(sec11_tbl, 1, "down")
    assert "【测试节顶部说明行】" in sec11_tbl.rows[2].cells[0].text

    # Move note row up
    move_row(sec11_tbl, 2, "up")
    assert "【测试节顶部说明行】" in sec11_tbl.rows[1].cells[0].text

        # Verify standardized export docx naming: {model} msds_{lang} {entity}.docx
    os_name = build_export_docx_name(sec11_doc, Path("OS-1030.docx"))
    assert os_name == "OS-1030 msds_CN 冠志.docx", f"Bad name: {os_name}"
    
    print("SELF_CHECK_PASS")


if __name__ == "__main__":
    main()
