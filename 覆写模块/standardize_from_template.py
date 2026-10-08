"""Standardize an MSDS source document with the formal Guanzhi template.

The source supplies every section row and every text value. The formal
template supplies row XML, merged-cell geometry, and formatting. Source files
are read-only; output rows are cloned from template rows and then receive only
source text.
"""

from __future__ import annotations

import copy
import re
import shutil
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from msds_template_editor import (
    all_runs,
    bold_runs,
    cell_views,
    looks_like_unbold_label,
    row_cells_with_spans,
)


@dataclass
class RowInfo:
    views: list
    labels: tuple[str, ...]
    roles: tuple[str, ...]
    spans: tuple[int, ...]
    paragraph_counts: tuple[int, ...]
    is_note: bool
    has_vertical_merge: bool


def label_key(text: str) -> str:
    text = re.sub(r"^\s*\d+(?:\.\d+)*[.\s]*", "", text or "")
    text = text.replace("\t", "").replace("\n", "")
    text = re.sub(r"[\s\u3000:：,，。；;（）()\[\]【】—–\-_/\\]", "", text)
    return text.lower()


def row_info(table_index: int, row_index: int, row) -> RowInfo:
    views = cell_views(table_index, row_index, row)
    spans = tuple(view.span for view in views)
    paragraph_counts = tuple(len(view.cell.paragraphs) for view in views)
    roles = []
    labels = []
    for view in views:
        if view.label_field and view.value_field:
            role = "label-value"
        elif view.label_field:
            role = "label"
        elif view.value_field:
            role = "note" if len(views) == 1 else "value"
        else:
            role = "structure"
        roles.append(role)
        if view.label_field:
            labels.append(label_key(view.label_field.value))
    has_vertical_merge = row._tr.find(".//" + qn("w:vMerge")) is not None
    return RowInfo(views, tuple(labels), tuple(roles), spans, paragraph_counts,
                   len(views) == 1 and roles == ["note"], has_vertical_merge)


def ensure_empty_value_runs(row):
    """Make empty cells writable without adding formatting properties."""
    spans = row_cells_with_spans(row)
    label_positions = [
        start
        for cell, start, _span in spans
        if bold_runs(cell) or looks_like_unbold_label(cell, len(spans))
    ]
    if not label_positions:
        return
    first_label = min(label_positions)
    for cell, start, _span in spans:
        if start > first_label and not all_runs(cell):
            cell.paragraphs[0].add_run("")


def ensure_role_runs(target_view, source_view):
    cell = target_view.cell
    if (source_view.label_field or source_view.value_field) and not target_view.label_field and not target_view.value_field:
        cell.paragraphs[0].add_run("")


def pair_views(source_views, target_views):
    if len(source_views) == len(target_views):
        return list(zip(source_views, target_views))
    if len(source_views) == 1 and target_views:
        target = next((view for view in reversed(target_views) if view.value_field), target_views[-1])
        return [(source_views[0], target)]
    if len(target_views) == 1 and source_views:
        return [(source_view, target_views[0]) for source_view in source_views]
    return list(zip(source_views, target_views))


def set_field(field, value):
    runs = list(field.runs)
    for run in runs:
        for child in list(run._r):
            if child.tag != qn("w:rPr"):
                run._r.remove(child)
    if runs:
        paragraphs = []
        for run in runs:
            parent_id = id(run._r.getparent())
            if not paragraphs or paragraphs[-1][0] != parent_id:
                paragraphs.append([parent_id, []])
            paragraphs[-1][1].append(run)
        parts = (value or "").split("\n")
        if len(parts) <= len(paragraphs):
            for index, (_parent_id, paragraph_runs) in enumerate(paragraphs):
                paragraph_runs[0].text = parts[index] if index < len(parts) else ""
                for run in paragraph_runs[1:]:
                    run.text = ""
        else:
            runs[0].text = value or ""
            for run in runs[1:]:
                run.text = ""
    field.value = value or ""


def source_cell_text(source_view):
    return source_view.cell.text


def copy_row_content(source_info: RowInfo, target_row, table_index: int, row_index: int):
    ensure_empty_value_runs(target_row)
    target_info = row_info(table_index, row_index, target_row)
    for view in target_info.views:
        if view.label_field:
            set_field(view.label_field, "")
        if view.value_field:
            set_field(view.value_field, "")
    if len(target_info.views) == 1 and len(source_info.views) > 1:
        text_parts = [source_cell_text(source_view) for source_view in source_info.views]
        target_view = target_info.views[0]
        target_field = target_view.value_field or target_view.label_field
        if target_field:
            set_field(target_field, "\n".join(part for part in text_parts if part))
        return
    pairs = pair_views(source_info.views, target_info.views)
    for source_view, target_view in pairs:
        ensure_role_runs(target_view, source_view)
    target_info = row_info(table_index, row_index, target_row)
    pairs = pair_views(source_info.views, target_info.views)
    for source_view, target_view in pairs:
        payload = source_cell_text(source_view)
        if source_view.label_field and source_view.value_field:
            if target_view.value_field:
                set_field(target_view.value_field, payload)
            elif target_view.label_field:
                set_field(target_view.label_field, payload)
        elif source_view.label_field:
            if target_view.label_field:
                set_field(target_view.label_field, payload)
            elif target_view.value_field:
                set_field(target_view.value_field, payload)
        elif source_view.value_field:
            if target_view.value_field:
                set_field(target_view.value_field, payload)
            elif target_view.label_field:
                set_field(target_view.label_field, payload)


def row_score(source: RowInfo, target: RowInfo) -> int:
    score = 12 * len(set(source.labels) & set(target.labels))
    score += 4 if source.spans == target.spans else 0
    score += 3 if source.roles == target.roles else 0
    score += 2 if len(source.views) == len(target.views) else 0
    score += 3 if source.is_note == target.is_note else 0
    if source.paragraph_counts == target.paragraph_counts:
        score += 12
    else:
        score -= 8 * abs(sum(source.paragraph_counts) - sum(target.paragraph_counts))
    score += 10 if not source.labels and target.is_note else 0
    score += 1 if bool(source.labels) == bool(target.labels) else 0
    return score


def choose_prototype(source_info: RowInfo, candidates: list[RowInfo]):
    if not source_info.has_vertical_merge:
        flat = [candidate for candidate in candidates if not candidate.has_vertical_merge]
        candidates = flat or candidates
    return max(candidates, key=lambda item: row_score(source_info, item))


def remove_all_rows(table):
    for row in list(table.rows):
        table._tbl.remove(row._tr)


def standardize_table(source_table, template_table, table_index):
    source_infos = [row_info(table_index, i, row) for i, row in enumerate(source_table.rows)]
    for row in template_table.rows:
        ensure_empty_value_runs(row)
    template_infos = [row_info(table_index, i, row) for i, row in enumerate(template_table.rows)]
    prototypes = template_infos[1:] or template_infos
    prototype_xml = {
        id(info): copy.deepcopy(template_table.rows[i]._tr)
        for i, info in enumerate(template_infos)
    }
    selected = []
    for source_index, source_info in enumerate(source_infos):
        prototype = template_infos[0] if source_index == 0 and template_infos else choose_prototype(source_info, prototypes)
        selected.append((source_info, prototype_xml[id(prototype)]))

    remove_all_rows(template_table)
    for output_index, (source_info, xml) in enumerate(selected):
        template_table._tbl.append(copy.deepcopy(xml))
        output_row = template_table.rows[-1]
        copy_row_content(source_info, output_row, table_index, output_index)


def standardize(source_path: Path, template_path: Path, output_path: Path):
    source = Document(str(source_path))
    work_dir = Path(tempfile.mkdtemp(prefix="msds-standardize-"))
    work_template = work_dir / template_path.name
    shutil.copy2(template_path, work_template)
    output_doc = Document(str(work_template))
    if len(source.tables) != len(output_doc.tables):
        raise ValueError(f"源文件与正式模板章节数不一致：{len(source.tables)} != {len(output_doc.tables)}")
    for table_index, source_table in enumerate(source.tables):
        standardize_table(source_table, output_doc.tables[table_index], table_index)
    output_doc.save(str(work_template))
    shutil.copy2(work_template, output_path)
    return output_path


def main(argv=None):
    argv = argv or sys.argv[1:]
    if len(argv) != 3:
        raise SystemExit("用法：python standardize_from_template.py 源文件.docx 正式模板.docx 输出文件.docx")
    output = standardize(Path(argv[0]).resolve(), Path(argv[1]).resolve(), Path(argv[2]).resolve())
    print(output)


if __name__ == "__main__":
    main()
